from io import BytesIO

from openpyxl import Workbook
from django.test import TestCase

from core.models import Alternativa, AlternativaCriterio, Criterio, CriterioParametro, Projeto
from core.services.spreadsheet_upload_service import prever_upload_planilha


class SpreadsheetUploadServiceTests(TestCase):
    def setUp(self):
        self.projeto = Projeto.objects.create(
            nome="Projeto planilha",
            descricao="Descricao do projeto",
            qtde_classes=3,
            qtde_criterios=3,
            qtde_alternativas=2,
            qtde_decisores=1,
            lamb=0.75,
        )
        self.criterio_custo = Criterio.objects.create(
            projeto=self.projeto,
            nome="Custo total",
            numerico=True,
            monotonico=2,
        )
        self.criterio_impacto = Criterio.objects.create(
            projeto=self.projeto,
            nome="Impacto social",
            numerico=False,
            monotonico=1,
        )
        self.criterio_emissoes = Criterio.objects.create(
            projeto=self.projeto,
            nome="Emissoes",
            numerico=True,
            monotonico=1,
        )
        self.alternativa_solar = Alternativa.objects.create(
            projeto=self.projeto,
            nome="Solar",
        )
        self.alternativa_eolica = Alternativa.objects.create(
            projeto=self.projeto,
            nome="Eolica",
        )

    def _workbook_bytes(
        self,
        *,
        criterios=None,
        alternativas=None,
        parametros=None,
        remover_abas=(),
        cabecalho_alternativas=None,
    ):
        workbook = Workbook()
        instrucoes = workbook.active
        instrucoes.title = "Instrucoes"
        instrucoes.append([
            "A planilha serve apenas para criterios numericos.",
            "q, p e v sao opcionais.",
        ])

        criterios_sheet = workbook.create_sheet("Criterios")
        criterios_sheet.append(("criterio_id", "nome", "tipo", "direcao", "preenchimento"))
        criterios = criterios if criterios is not None else [
            (self.criterio_custo.id, "Custo total", "numerico", "custo", "preencher em Alternativas"),
            (self.criterio_impacto.id, "Impacto social", "qualitativo", "lucro", "avaliar no app"),
            (self.criterio_emissoes.id, "Emissoes", "numerico", "lucro", "preencher em Alternativas"),
        ]
        for linha in criterios:
            criterios_sheet.append(linha)

        alternativas_sheet = workbook.create_sheet("Alternativas")
        cabecalho_alternativas = cabecalho_alternativas or (
            "alternativa_id", "alternativa", "Custo total", "Emissoes")
        alternativas_sheet.append(cabecalho_alternativas)
        alternativas = alternativas if alternativas is not None else [
            (self.alternativa_solar.id, "Solar", 100, 30),
            (self.alternativa_eolica.id, "Eolica", 80, 20),
        ]
        for linha in alternativas:
            alternativas_sheet.append(linha)

        parametros_sheet = workbook.create_sheet("Parametros")
        parametros_sheet.append(("criterio_id", "criterio", "q", "p", "v"))
        parametros = parametros if parametros is not None else [
            (self.criterio_custo.id, "Custo total", None, None, None),
            (self.criterio_emissoes.id, "Emissoes", 2, None, None),
        ]
        for linha in parametros:
            parametros_sheet.append(linha)

        for nome_aba in remover_abas:
            del workbook[nome_aba]

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer

    def test_prever_upload_calcula_parametros_e_nao_grava_banco(self):
        self.assertEqual(AlternativaCriterio.objects.count(), 0)
        self.assertEqual(CriterioParametro.objects.count(), 0)

        preview = prever_upload_planilha(self.projeto, self._workbook_bytes())

        self.assertEqual(AlternativaCriterio.objects.count(), 0)
        self.assertEqual(CriterioParametro.objects.count(), 0)
        self.assertEqual(preview["erros"], [])
        self.assertGreaterEqual(len(preview["avisos"]), 1)

        criterios = {item["criterio_id"]: item for item in preview["parametros"]}
        custo = criterios[self.criterio_custo.id]
        emissoes = criterios[self.criterio_emissoes.id]

        self.assertEqual(custo["q"]["valor"], 4.0)
        self.assertEqual(custo["q"]["origem"], "automatico")
        self.assertEqual(custo["p"]["valor"], 8.0)
        self.assertEqual(custo["v"]["valor"], 18.0)
        self.assertEqual(emissoes["q"]["valor"], 2.0)
        self.assertEqual(emissoes["q"]["origem"], "planilha")
        self.assertEqual(emissoes["p"]["valor"], 4.0)
        self.assertEqual(emissoes["p"]["origem"], "automatico")
        self.assertEqual(emissoes["v"]["valor"], 9.0)

        desempenhos = {
            item["alternativa_id"]: item for item in preview["desempenhos"]
        }
        self.assertEqual(desempenhos[self.alternativa_solar.id]["valores"][0]["valor"], 100.0)

    def test_prever_upload_preserva_parametros_preenchidos_parcialmente(self):
        parametros = [
            (self.criterio_custo.id, "Custo total", 1, None, 30),
            (self.criterio_emissoes.id, "Emissoes", None, 5, None),
        ]

        preview = prever_upload_planilha(
            self.projeto,
            self._workbook_bytes(parametros=parametros),
        )

        criterios = {item["criterio_id"]: item for item in preview["parametros"]}
        custo = criterios[self.criterio_custo.id]
        emissoes = criterios[self.criterio_emissoes.id]

        self.assertEqual(custo["q"]["valor"], 1.0)
        self.assertEqual(custo["q"]["origem"], "planilha")
        self.assertEqual(custo["p"]["valor"], 8.0)
        self.assertEqual(custo["p"]["origem"], "automatico")
        self.assertEqual(custo["v"]["valor"], 30.0)
        self.assertEqual(custo["v"]["origem"], "planilha")
        self.assertEqual(emissoes["q"]["valor"], 2.0)
        self.assertEqual(emissoes["q"]["origem"], "automatico")
        self.assertEqual(emissoes["p"]["valor"], 5.0)
        self.assertEqual(emissoes["p"]["origem"], "planilha")
        self.assertEqual(emissoes["v"]["valor"], 9.0)
        self.assertEqual(emissoes["v"]["origem"], "automatico")

    def test_prever_upload_ignora_criterios_qualitativos_com_aviso(self):
        preview = prever_upload_planilha(self.projeto, self._workbook_bytes())

        avisos = " ".join(aviso["codigo"] for aviso in preview["avisos"])
        self.assertIn("criterio_qualitativo_ignorado", avisos)
        self.assertTrue(
            any(
                aviso["codigo"] == "criterio_qualitativo_ignorado"
                and "Impacto social" in aviso["mensagem"]
                for aviso in preview["avisos"]
            )
        )

    def test_prever_upload_rejeita_abas_obrigatorias_ausentes(self):
        preview = prever_upload_planilha(
            self.projeto,
            self._workbook_bytes(remover_abas=("Criterios", "Alternativas", "Parametros")),
        )

        self.assertTrue(preview["erros"])
        self.assertTrue(
            any(erro["codigo"] == "aba_obrigatoria_ausente" for erro in preview["erros"])
        )

    def test_prever_upload_rejeita_criterio_desconhecido_e_desempenho_vazio(self):
        alternativas = [
            (self.alternativa_solar.id, "Solar", 100, 30, 999),
            (self.alternativa_eolica.id, "Eolica", None, 20, None),
        ]
        workbook = self._workbook_bytes(
            alternativas=alternativas,
            cabecalho_alternativas=(
                "alternativa_id",
                "alternativa",
                "Custo total",
                "Emissoes",
                "Desconhecido",
            ),
        )

        preview = prever_upload_planilha(self.projeto, workbook)

        codigos = {erro["codigo"] for erro in preview["erros"]}
        self.assertIn("criterio_desconhecido", codigos)
        self.assertIn("desempenho_vazio", codigos)

    def test_prever_upload_rejeita_alternativa_desconhecida_e_numero_invalido(self):
        alternativas = [
            (self.alternativa_solar.id, "Solar", "100,5", 30),
            (999, "Outra", 80, "texto"),
        ]
        workbook = self._workbook_bytes(alternativas=alternativas)

        preview = prever_upload_planilha(self.projeto, workbook)

        codigos = {erro["codigo"] for erro in preview["erros"]}
        self.assertIn("alternativa_desconhecida", codigos)
        self.assertIn("numero_invalido", codigos)
