from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook

from core.models import Alternativa, Criterio, Projeto
from core.services.spreadsheet_template_service import gerar_planilha_modelo


class SpreadsheetTemplateServiceTests(TestCase):
    def setUp(self):
        self.projeto = Projeto.objects.create(
            nome="Projeto planilha",
            descricao="Descricao do projeto",
            qtde_classes=3,
            qtde_criterios=3,
            qtde_alternativas=2,
            qtde_decisores=1,
            lamb=0.75,
        )
        self.criterio_numerico_1 = Criterio.objects.create(
            projeto=self.projeto,
            nome="Custo total",
            numerico=True,
            monotonico=2,
        )
        self.criterio_qualitativo = Criterio.objects.create(
            projeto=self.projeto,
            nome="Impacto social",
            numerico=False,
            monotonico=1,
        )
        self.criterio_numerico_2 = Criterio.objects.create(
            projeto=self.projeto,
            nome="Emissoes",
            numerico=True,
            monotonico=1,
        )
        self.alternativa_1 = Alternativa.objects.create(
            projeto=self.projeto,
            nome="Solar",
        )
        self.alternativa_2 = Alternativa.objects.create(
            projeto=self.projeto,
            nome="Eolica",
        )

    def test_gera_workbook_xlsx_com_abas_e_conteudo_esperados(self):
        conteudo_xlsx = gerar_planilha_modelo(self.projeto)

        workbook = load_workbook(BytesIO(conteudo_xlsx))

        self.assertEqual(
            workbook.sheetnames,
            ["Instrucoes", "Criterios", "Alternativas", "Parametros"],
        )

        instrucoes = workbook["Instrucoes"]
        texto_instrucoes = "\n".join(
            str(celula)
            for linha in instrucoes.iter_rows(values_only=True)
            for celula in linha
            if celula is not None
        )
        self.assertIn("criterios numericos", texto_instrucoes.lower())
        self.assertIn("q, p e v", texto_instrucoes.lower())
        self.assertIn("20%", texto_instrucoes)
        self.assertIn("40%", texto_instrucoes)
        self.assertIn("90%", texto_instrucoes)

        criterios = workbook["Criterios"]
        linhas_criterios = list(criterios.iter_rows(values_only=True))
        self.assertEqual(
            linhas_criterios[0],
            ("criterio_id", "nome", "tipo", "direcao", "preenchimento"),
        )
        self.assertEqual(
            linhas_criterios[1],
            (
                self.criterio_numerico_1.id,
                "Custo total",
                "numerico",
                "custo",
                "preencher em Alternativas",
            ),
        )
        self.assertEqual(
            linhas_criterios[2],
            (
                self.criterio_qualitativo.id,
                "Impacto social",
                "qualitativo",
                "lucro",
                "avaliar no app",
            ),
        )
        self.assertEqual(
            linhas_criterios[3],
            (
                self.criterio_numerico_2.id,
                "Emissoes",
                "numerico",
                "lucro",
                "preencher em Alternativas",
            ),
        )

        alternativas = workbook["Alternativas"]
        linhas_alternativas = list(alternativas.iter_rows(values_only=True))
        self.assertEqual(
            linhas_alternativas[0],
            (
                "alternativa_id",
                "alternativa",
                "Custo total",
                "Emissoes",
            ),
        )
        self.assertEqual(
            linhas_alternativas[1],
            (self.alternativa_1.id, "Solar", None, None),
        )
        self.assertEqual(
            linhas_alternativas[2],
            (self.alternativa_2.id, "Eolica", None, None),
        )
        self.assertNotIn("Impacto social", linhas_alternativas[0])

        parametros = workbook["Parametros"]
        linhas_parametros = list(parametros.iter_rows(values_only=True))
        self.assertEqual(
            linhas_parametros[0],
            ("criterio_id", "criterio", "q", "p", "v"),
        )
        self.assertEqual(
            linhas_parametros[1],
            (self.criterio_numerico_1.id, "Custo total", None, None, None),
        )
        self.assertEqual(
            linhas_parametros[2],
            (self.criterio_numerico_2.id, "Emissoes", None, None, None),
        )
