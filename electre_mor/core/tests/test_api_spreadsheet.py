from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from core.models import Alternativa, AlternativaCriterio, Criterio, CriterioParametro, Decisor, Projeto


SPREADSHEET_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


class ApiSpreadsheetTests(APITestCase):
    def setUp(self):
        self.projeto = Projeto.objects.create(
            nome="Projeto planilha",
            descricao="Descricao do projeto",
            qtde_classes=3,
            qtde_criterios=2,
            qtde_alternativas=2,
            qtde_decisores=1,
            lamb=0.75,
        )
        self.criterio_custo = Criterio.objects.create(
            projeto=self.projeto,
            nome="Custo total",
            numerico=True,
            monotonico=2,
        )
        self.criterio_emissoes = Criterio.objects.create(
            projeto=self.projeto,
            nome="Emissoes",
            numerico=True,
            monotonico=1,
        )
        self.alternativa_solar = Alternativa.objects.create(
            projeto=self.projeto,
            nome="Solar",
        )
        self.alternativa_eolica = Alternativa.objects.create(
            projeto=self.projeto,
            nome="Eolica",
        )

    def _criar_workbook_valido(self) -> bytes:
        workbook = Workbook()
        instrucoes = workbook.active
        instrucoes.title = "Instrucoes"
        instrucoes.append(["Texto", "Outro texto"])

        criterios = workbook.create_sheet("Criterios")
        criterios.append(("criterio_id", "nome", "tipo", "direcao", "preenchimento"))
        criterios.append(
            (self.criterio_custo.id, self.criterio_custo.nome, "numerico", "custo", "preencher em Alternativas")
        )
        criterios.append(
            (
                self.criterio_emissoes.id,
                self.criterio_emissoes.nome,
                "numerico",
                "lucro",
                "preencher em Alternativas",
            )
        )

        alternativas = workbook.create_sheet("Alternativas")
        alternativas.append(("alternativa_id", "alternativa", self.criterio_custo.nome, self.criterio_emissoes.nome))
        alternativas.append((self.alternativa_solar.id, self.alternativa_solar.nome, 100, 30))
        alternativas.append((self.alternativa_eolica.id, self.alternativa_eolica.nome, 80, 20))

        parametros = workbook.create_sheet("Parametros")
        parametros.append(("criterio_id", "criterio", "q", "p", "v"))
        parametros.append((self.criterio_custo.id, self.criterio_custo.nome, None, None, None))
        parametros.append((self.criterio_emissoes.id, self.criterio_emissoes.nome, 2, 4, 9))

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def _criar_upload(self, nome: str, conteudo: bytes) -> SimpleUploadedFile:
        return SimpleUploadedFile(nome, conteudo, content_type=SPREADSHEET_CONTENT_TYPE)

    def _payload_confirmacao(self):
        return {
            "desempenhos": [
                {
                    "alternativa_id": self.alternativa_solar.id,
                    "alternativa": self.alternativa_solar.nome,
                    "valores": [
                        {
                            "criterio_id": self.criterio_custo.id,
                            "criterio": self.criterio_custo.nome,
                            "valor": 100,
                        },
                        {
                            "criterio_id": self.criterio_emissoes.id,
                            "criterio": self.criterio_emissoes.nome,
                            "valor": 30,
                        },
                    ],
                },
                {
                    "alternativa_id": self.alternativa_eolica.id,
                    "alternativa": self.alternativa_eolica.nome,
                    "valores": [
                        {
                            "criterio_id": self.criterio_custo.id,
                            "criterio": self.criterio_custo.nome,
                            "valor": 80,
                        },
                        {
                            "criterio_id": self.criterio_emissoes.id,
                            "criterio": self.criterio_emissoes.nome,
                            "valor": 20,
                        },
                    ],
                },
            ],
            "parametros": [
                {
                    "criterio_id": self.criterio_custo.id,
                    "criterio": self.criterio_custo.nome,
                    "q": {"valor": 2, "origem": "automatico"},
                    "p": {"valor": 4, "origem": "automatico"},
                    "v": {"valor": 9, "origem": "automatico"},
                },
                {
                    "criterio_id": self.criterio_emissoes.id,
                    "criterio": self.criterio_emissoes.nome,
                    "q": {"valor": 1, "origem": "editado"},
                    "p": {"valor": 2, "origem": "planilha"},
                    "v": {"valor": 3, "origem": "editado"},
                },
            ],
        }

    def test_download_spreadsheet_template_retorna_xlsx(self):
        response = self.client.get(
            f"/api/v1/projects/{self.projeto.id}/spreadsheet-template/"
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], SPREADSHEET_CONTENT_TYPE)

        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(
            workbook.sheetnames,
            ["Instrucoes", "Criterios", "Alternativas", "Parametros"],
        )

    def test_preview_spreadsheet_multipart_retorna_resposta_estruturada(self):
        upload = self._criar_upload("modelo.xlsx", self._criar_workbook_valido())

        response = self.client.post(
            f"/api/v1/projects/{self.projeto.id}/spreadsheet-upload/preview/",
            {"arquivo": upload},
            format="multipart",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["projeto"]["id"], self.projeto.id)
        self.assertEqual(response.data["arquivo"]["nome"], "modelo.xlsx")
        self.assertEqual(response.data["arquivo"]["tipo"], SPREADSHEET_CONTENT_TYPE)
        self.assertEqual(len(response.data["desempenhos"]), 2)
        self.assertEqual(len(response.data["parametros"]), 2)
        self.assertEqual(response.data["erros"], [])
        self.assertEqual(AlternativaCriterio.objects.count(), 0)
        self.assertEqual(CriterioParametro.objects.count(), 0)

    def test_preview_spreadsheet_com_erros_retorna_200(self):
        workbook = Workbook()
        workbook.active.title = "Instrucoes"
        workbook.create_sheet("Criterios")
        workbook.create_sheet("Alternativas")
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        response = self.client.post(
            f"/api/v1/projects/{self.projeto.id}/spreadsheet-upload/preview/",
            {"arquivo": self._criar_upload("incompleta.xlsx", buffer.getvalue())},
            format="multipart",
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data["erros"])
        self.assertGreater(len(response.data["erros"]), 0)

    def test_confirm_spreadsheet_salva_notas_e_parametros_para_criador(self):
        criador = Decisor.objects.create(
            projeto=self.projeto,
            nome="Criador",
            is_criador=True,
        )

        response = self.client.post(
            f"/api/v1/projects/{self.projeto.id}/spreadsheet-upload/confirm/",
            self._payload_confirmacao(),
            format="json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["projeto"]["id"], self.projeto.id)
        self.assertEqual(
            AlternativaCriterio.objects.filter(projeto=self.projeto, decisor=criador).count(),
            4,
        )
        self.assertEqual(
            CriterioParametro.objects.filter(projeto=self.projeto).count(),
            2,
        )
        self.assertEqual(
            list(
                AlternativaCriterio.objects.filter(
                    projeto=self.projeto,
                    decisor=criador,
                ).order_by("criterio_id", "alternativa_id").values_list("criterio_id", "alternativa_id", "nota")
            ),
            [
                (self.criterio_custo.id, self.alternativa_solar.id, 100.0),
                (self.criterio_custo.id, self.alternativa_eolica.id, 80.0),
                (self.criterio_emissoes.id, self.alternativa_solar.id, 30.0),
                (self.criterio_emissoes.id, self.alternativa_eolica.id, 20.0),
            ],
        )

    def test_confirm_spreadsheet_sem_criador_usa_primeiro_decisor_compativel(self):
        decisor = Decisor.objects.create(
            projeto=self.projeto,
            nome="Decisor unico",
            is_criador=False,
        )

        response = self.client.post(
            f"/api/v1/projects/{self.projeto.id}/spreadsheet-upload/confirm/",
            self._payload_confirmacao(),
            format="json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            AlternativaCriterio.objects.filter(projeto=self.projeto, decisor=decisor).count(),
            4,
        )
        self.assertEqual(
            AlternativaCriterio.objects.exclude(decisor=decisor).count(),
            0,
        )

    def test_confirm_spreadsheet_bloqueia_quando_resultado_ja_foi_gerado(self):
        Decisor.objects.create(
            projeto=self.projeto,
            nome="Criador",
            is_criador=True,
        )
        self.projeto.resultado_gerado_em = timezone.now()
        self.projeto.save(update_fields=["resultado_gerado_em"])

        response = self.client.post(
            f"/api/v1/projects/{self.projeto.id}/spreadsheet-upload/confirm/",
            self._payload_confirmacao(),
            format="json",
        )

        self.assertEqual(response.status_code, 409)
        self.assertEqual(AlternativaCriterio.objects.count(), 0)
        self.assertEqual(CriterioParametro.objects.count(), 0)
