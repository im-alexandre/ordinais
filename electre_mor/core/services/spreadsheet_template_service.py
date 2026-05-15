from __future__ import annotations

from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from core.models import Criterio, Projeto


def _ordenar_criterios(projeto: Projeto) -> list[Criterio]:
    return list(projeto.criterios.all().order_by("id"))


def _ordenar_alternativas(projeto: Projeto):
    return list(projeto.alternativas.all().order_by("id"))


def _direcao_texto(criterio: Criterio) -> str:
    return "lucro" if criterio.monotonico == 1 else "custo"


def _tipo_texto(criterio: Criterio) -> str:
    return "numerico" if criterio.numerico else "qualitativo"


def _preenchimento_texto(criterio: Criterio) -> str:
    return "preencher em Alternativas" if criterio.numerico else "avaliar no app"


def _formatar_cabecalho(worksheet: Worksheet) -> None:
    for cell in next(worksheet.iter_rows(min_row=1, max_row=1)):
        cell.font = Font(bold=True)


def _criar_aba_instrucoes(workbook: Workbook) -> Worksheet:
    worksheet = workbook.create_sheet("Instrucoes")
    instrucoes = [
        [
            "A planilha serve apenas para criterios numericos.",
            "Criterios qualitativos continuam sendo avaliados no app.",
        ],
        [
            "Campos q, p e v sao opcionais.",
            "Se ficarem vazios, o sistema calcula automaticamente q = 20%, p = 40% e v = 90% da maior diferenca de desempenho do criterio.",
        ],
        [
            "Nao alterar nomes de abas, nomes de criterios ou nomes de alternativas.",
        ],
    ]
    for linha in instrucoes:
        worksheet.append(linha)
    worksheet.column_dimensions["A"].width = 80
    worksheet.column_dimensions["B"].width = 110
    worksheet.freeze_panes = "A1"
    return worksheet


def _criar_aba_criterios(workbook: Workbook, projeto: Projeto,
                         criterios: Iterable[Criterio]) -> Worksheet:
    worksheet = workbook.create_sheet("Criterios")
    worksheet.append(
        ("criterio_id", "nome", "tipo", "direcao", "preenchimento"))
    for criterio in criterios:
        worksheet.append((
            criterio.id,
            criterio.nome,
            _tipo_texto(criterio),
            _direcao_texto(criterio),
            _preenchimento_texto(criterio),
        ))
    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 14
    worksheet.column_dimensions["B"].width = 24
    worksheet.column_dimensions["C"].width = 14
    worksheet.column_dimensions["D"].width = 14
    worksheet.column_dimensions["E"].width = 28
    return worksheet


def _criar_aba_alternativas(workbook: Workbook, projeto: Projeto,
                            criterios: list[Criterio]) -> Worksheet:
    worksheet = workbook.create_sheet("Alternativas")
    criterios_numericos = [criterio for criterio in criterios if criterio.numerico]
    cabecalho = ["alternativa_id", "alternativa"]
    cabecalho.extend(criterio.nome for criterio in criterios_numericos)
    worksheet.append(tuple(cabecalho))

    for alternativa in _ordenar_alternativas(projeto):
        linha = [alternativa.id, alternativa.nome]
        linha.extend([None] * len(criterios_numericos))
        worksheet.append(tuple(linha))

    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 16
    worksheet.column_dimensions["B"].width = 24
    for indice in range(3, len(cabecalho) + 1):
        worksheet.column_dimensions[get_column_letter(indice)].width = 18
    return worksheet


def _criar_aba_parametros(workbook: Workbook, criterios: list[Criterio]) -> Worksheet:
    worksheet = workbook.create_sheet("Parametros")
    worksheet.append(("criterio_id", "criterio", "q", "p", "v"))
    for criterio in criterios:
        if not criterio.numerico:
            continue
        worksheet.append((criterio.id, criterio.nome, None, None, None))
    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 14
    worksheet.column_dimensions["B"].width = 24
    worksheet.column_dimensions["C"].width = 12
    worksheet.column_dimensions["D"].width = 12
    worksheet.column_dimensions["E"].width = 12
    return worksheet


def gerar_planilha_modelo(projeto: Projeto) -> bytes:
    """Gera a planilha modelo XLSX do projeto configurado."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    criterios = _ordenar_criterios(projeto)
    _criar_aba_instrucoes(workbook)
    _criar_aba_criterios(workbook, projeto, criterios)
    _criar_aba_alternativas(workbook, projeto, criterios)
    _criar_aba_parametros(workbook, criterios)

    for worksheet in workbook.worksheets:
        _formatar_cabecalho(worksheet)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
