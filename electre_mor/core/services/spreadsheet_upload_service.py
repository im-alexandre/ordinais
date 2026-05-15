from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from core.models import Criterio, Projeto


ABAS_OBRIGATORIAS = ("Instrucoes", "Criterios", "Alternativas", "Parametros")


def _erro(codigo: str, mensagem: str, **contexto: Any) -> dict[str, Any]:
    payload = {"codigo": codigo, "mensagem": mensagem}
    if contexto:
        payload.update(contexto)
    return payload


def _aviso(codigo: str, mensagem: str, **contexto: Any) -> dict[str, Any]:
    payload = {"codigo": codigo, "mensagem": mensagem}
    if contexto:
        payload.update(contexto)
    return payload


def _normalizar_texto(valor: Any) -> str:
    if valor is None:
        return ""
    return str(valor).strip()


def _valor_vazio(valor: Any) -> bool:
    return valor is None or _normalizar_texto(valor) == ""


def _normalizar_numero(valor: Any) -> Decimal | None:
    if _valor_vazio(valor):
        return None
    if isinstance(valor, bool):
        raise ValueError("bool nao e numero")
    if isinstance(valor, Decimal):
        return valor
    if isinstance(valor, (int, float)):
        return Decimal(str(valor))

    texto = _normalizar_texto(valor).replace(" ", "")
    if texto == "":
        return None

    try:
        if "," in texto and "." in texto:
            ultimo = max(texto.rfind(","), texto.rfind("."))
            inteiro = texto[:ultimo].replace(",", "").replace(".", "")
            fracao = texto[ultimo + 1 :]
            texto = f"{inteiro}.{fracao}"
        elif "," in texto:
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
        return Decimal(texto)
    except (InvalidOperation, ValueError):
        raise ValueError("numero invalido") from None


def _ler_workbook(arquivo: Any):
    if hasattr(arquivo, "seek"):
        arquivo.seek(0)
    if isinstance(arquivo, (bytes, bytearray)):
        return load_workbook(BytesIO(arquivo), data_only=True)
    return load_workbook(arquivo, data_only=True)


def _criterios_do_projeto(projeto: Projeto) -> list[Criterio]:
    return list(projeto.criterios.all().order_by("id"))


def _alternativas_do_projeto(projeto: Projeto):
    return list(projeto.alternativas.all().order_by("id"))


def _identificar_aba_obrigatoria(workbook) -> list[dict[str, Any]]:
    erros: list[dict[str, Any]] = []
    presentes = set(workbook.sheetnames)
    for aba in ABAS_OBRIGATORIAS:
        if aba not in presentes:
            erros.append(
                _erro(
                    "aba_obrigatoria_ausente",
                    f"Aba obrigatoria ausente: {aba}.",
                    aba=aba,
                )
            )
    return erros


def _row_iter(worksheet):
    for linha in worksheet.iter_rows(values_only=True):
        yield linha


def _sheet_rows(worksheet) -> list[tuple[Any, ...]]:
    linhas = list(_row_iter(worksheet))
    if not linhas:
        return []
    return [tuple(linha) for linha in linhas[1:]]


def _validar_criterios(
    projeto: Projeto,
    workbook,
) -> tuple[list[Criterio], list[dict[str, Any]]]:
    worksheet = workbook["Criterios"]
    criterios_projeto = _criterios_do_projeto(projeto)
    criterios_por_id = {str(criterio.id): criterio for criterio in criterios_projeto}
    criterios_por_nome = {criterio.nome: criterio for criterio in criterios_projeto}
    avisos: list[dict[str, Any]] = []
    erros: list[dict[str, Any]] = []
    criterios_sheet_ids: set[str] = set()
    criterios_sheet_nomes: set[str] = set()

    for linha in _sheet_rows(worksheet):
        criterio_id = _normalizar_texto(linha[0] if len(linha) > 0 else None)
        criterio_nome = _normalizar_texto(linha[1] if len(linha) > 1 else None)
        criterio = criterios_por_id.get(criterio_id) or criterios_por_nome.get(criterio_nome)
        if criterio is None:
            erros.append(
                _erro(
                    "criterio_desconhecido",
                    f"Criterio desconhecido: {criterio_nome or criterio_id}.",
                    criterio_id=criterio_id,
                    criterio=criterio_nome,
                )
            )
            continue
        criterios_sheet_ids.add(str(criterio.id))
        criterios_sheet_nomes.add(criterio.nome)
        if not criterio.numerico:
            avisos.append(
                _aviso(
                    "criterio_qualitativo_ignorado",
                    f"Criterio qualitativo ignorado na planilha: {criterio.nome}.",
                    criterio_id=criterio.id,
                    criterio=criterio.nome,
                )
            )

    criterios_numericos = [criterio for criterio in criterios_projeto if criterio.numerico]
    for criterio in criterios_numericos:
        if str(criterio.id) not in criterios_sheet_ids and criterio.nome not in criterios_sheet_nomes:
            erros.append(
                _erro(
                    "criterio_numerico_ausente",
                    f"Criterio numerico ausente na aba Criterios: {criterio.nome}.",
                    criterio_id=criterio.id,
                    criterio=criterio.nome,
                )
            )

    return criterios_numericos, avisos + erros


def _mapear_colunas_alternativas(
    worksheet,
    criterios_numericos: list[Criterio],
) -> tuple[dict[int, Criterio], list[dict[str, Any]], list[dict[str, Any]]]:
    cabecalho = next(_row_iter(worksheet))
    criterios_por_nome = {criterio.nome: criterio for criterio in criterios_numericos}
    criterios_por_id = {str(criterio.id): criterio for criterio in criterios_numericos}
    mapeamento: dict[int, Criterio] = {}
    avisos: list[dict[str, Any]] = []
    erros: list[dict[str, Any]] = []

    for indice, nome_coluna in enumerate(cabecalho[2:], start=2):
        texto_coluna = _normalizar_texto(nome_coluna)
        if texto_coluna == "":
            continue
        criterio = criterios_por_nome.get(texto_coluna) or criterios_por_id.get(texto_coluna)
        if criterio is None:
            valores_coluna = [
                linha[indice] if len(linha) > indice else None
                for linha in _sheet_rows(worksheet)
            ]
            if any(not _valor_vazio(valor) for valor in valores_coluna):
                erros.append(
                    _erro(
                        "criterio_desconhecido",
                        f"Coluna de criterio desconhecida: {texto_coluna}.",
                        coluna=texto_coluna,
                    )
                )
            else:
                avisos.append(
                    _aviso(
                        "coluna_extra_ignorada",
                        f"Coluna extra ignorada: {texto_coluna}.",
                        coluna=texto_coluna,
                    )
                )
            continue
        mapeamento[indice] = criterio
    return mapeamento, avisos, erros


def _validar_alternativas(
    projeto: Projeto,
    workbook,
    criterios_numericos: list[Criterio],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    worksheet = workbook["Alternativas"]
    alternativas_projeto = _alternativas_do_projeto(projeto)
    alternativas_por_id = {str(alternativa.id): alternativa for alternativa in alternativas_projeto}
    alternativas_por_nome = {alternativa.nome: alternativa for alternativa in alternativas_projeto}

    mapeamento, avisos, erros = _mapear_colunas_alternativas(worksheet, criterios_numericos)
    criterios_presentes = {criterio.id for criterio in mapeamento.values()}
    for criterio in criterios_numericos:
        if criterio.id not in criterios_presentes:
            erros.append(
                _erro(
                    "criterio_numerico_ausente",
                    f"Criterio numerico ausente na aba Alternativas: {criterio.nome}.",
                    criterio_id=criterio.id,
                    criterio=criterio.nome,
                )
            )

    desempenhos: list[dict[str, Any]] = []
    alternativas_presentes: set[int] = set()

    for linha in _sheet_rows(worksheet):
        alternativa_id = _normalizar_texto(linha[0] if len(linha) > 0 else None)
        alternativa_nome = _normalizar_texto(linha[1] if len(linha) > 1 else None)
        alternativa = alternativas_por_id.get(alternativa_id) or alternativas_por_nome.get(alternativa_nome)
        row_result = {
            "alternativa_id": int(alternativa.id) if alternativa is not None else None,
            "alternativa": alternativa.nome if alternativa is not None else alternativa_nome or alternativa_id,
            "valores": [],
        }

        if alternativa is None:
            erros.append(
                _erro(
                    "alternativa_desconhecida",
                    f"Alternativa desconhecida: {alternativa_nome or alternativa_id}.",
                    alternativa_id=alternativa_id,
                    alternativa=alternativa_nome,
                )
            )
        else:
            alternativas_presentes.add(alternativa.id)

        for indice, criterio in mapeamento.items():
            valor_bruto = linha[indice] if len(linha) > indice else None
            contexto = {
                "alternativa_id": alternativa.id if alternativa is not None else alternativa_id,
                "alternativa": alternativa.nome if alternativa is not None else alternativa_nome,
                "criterio_id": criterio.id,
                "criterio": criterio.nome,
            }
            if _valor_vazio(valor_bruto):
                erros.append(_erro("desempenho_vazio", "Desempenho numerico vazio.", **contexto))
                continue
            try:
                valor = _normalizar_numero(valor_bruto)
            except ValueError:
                erros.append(
                    _erro(
                        "numero_invalido",
                        "Valor numerico invalido.",
                        valor=valor_bruto,
                        **contexto,
                    )
                )
                continue
            if valor is None:
                erros.append(_erro("numero_invalido", "Valor numerico invalido.", **contexto))
                continue
            if alternativa is not None:
                row_result["valores"].append(
                    {
                        "criterio_id": criterio.id,
                        "criterio": criterio.nome,
                        "valor": float(valor),
                    }
                )

        desempenhos.append(row_result)

    for alternativa in alternativas_projeto:
        if alternativa.id not in alternativas_presentes:
            erros.append(
                _erro(
                    "alternativa_ausente",
                    f"Alternativa ausente na aba Alternativas: {alternativa.nome}.",
                    alternativa_id=alternativa.id,
                    alternativa=alternativa.nome,
                )
            )

    return desempenhos, avisos, erros


def _max_diferenca(valores: list[Decimal]) -> Decimal:
    if not valores:
        return Decimal("0")
    return max(valores) - min(valores)


def _resolver_parametro(valor_planilha: Any, diff: Decimal, proporcao: Decimal) -> tuple[float, str]:
    if _valor_vazio(valor_planilha):
        return float(diff * proporcao), "automatico"
    valor = _normalizar_numero(valor_planilha)
    if valor is None:
        return float(diff * proporcao), "automatico"
    return float(valor), "planilha"


def _validar_parametros(
    workbook,
    criterios_numericos: list[Criterio],
    desempenhos: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    worksheet = workbook["Parametros"]
    rows = _sheet_rows(worksheet)
    por_id = {str(row[0]): row for row in rows if len(row) > 0}
    por_nome = {_normalizar_texto(row[1]): row for row in rows if len(row) > 1}
    erros: list[dict[str, Any]] = []
    resultados: list[dict[str, Any]] = []

    desempenhos_por_criterio: dict[int, list[Decimal]] = {
        criterio.id: [] for criterio in criterios_numericos
    }
    for desempenho in desempenhos:
        for item in desempenho["valores"]:
            desempenhos_por_criterio[item["criterio_id"]].append(Decimal(str(item["valor"])))

    for criterio in criterios_numericos:
        row = por_id.get(str(criterio.id)) or por_nome.get(criterio.nome)
        if row is None:
            erros.append(
                _erro(
                    "criterio_parametro_ausente",
                    f"Parametro ausente para criterio numerico: {criterio.nome}.",
                    criterio_id=criterio.id,
                    criterio=criterio.nome,
                )
            )
            continue

        diff = _max_diferenca(desempenhos_por_criterio.get(criterio.id, []))
        q, origem_q = _resolver_parametro(row[2] if len(row) > 2 else None, diff, Decimal("0.20"))
        p, origem_p = _resolver_parametro(row[3] if len(row) > 3 else None, diff, Decimal("0.40"))
        v, origem_v = _resolver_parametro(row[4] if len(row) > 4 else None, diff, Decimal("0.90"))

        resultados.append(
            {
                "criterio_id": criterio.id,
                "criterio": criterio.nome,
                "q": {"valor": q, "origem": origem_q},
                "p": {"valor": p, "origem": origem_p},
                "v": {"valor": v, "origem": origem_v},
                "maior_diferenca": float(diff),
            }
        )

    return resultados, erros


def prever_upload_planilha(projeto: Projeto, arquivo: Any) -> dict[str, Any]:
    workbook = _ler_workbook(arquivo)
    erros = _identificar_aba_obrigatoria(workbook)
    avisos: list[dict[str, Any]] = []
    if erros:
        return {
            "arquivo": {"abas": workbook.sheetnames},
            "avisos": avisos,
            "erros": erros,
            "desempenhos": [],
            "parametros": [],
        }

    criterios_numericos, avisos_e_erros_criterios = _validar_criterios(projeto, workbook)
    avisos.extend([
        item for item in avisos_e_erros_criterios
        if item["codigo"] == "criterio_qualitativo_ignorado"
    ])
    erros.extend([
        item for item in avisos_e_erros_criterios
        if item["codigo"] != "criterio_qualitativo_ignorado"
    ])

    desempenhos, avisos_alternativas, erros_alternativas = _validar_alternativas(
        projeto,
        workbook,
        criterios_numericos,
    )
    avisos.extend(avisos_alternativas)
    erros.extend(erros_alternativas)

    parametros, erros_parametros = _validar_parametros(workbook, criterios_numericos, desempenhos)
    erros.extend(erros_parametros)

    return {
        "arquivo": {"abas": workbook.sheetnames},
        "avisos": avisos,
        "erros": erros,
        "desempenhos": desempenhos,
        "parametros": parametros,
    }
